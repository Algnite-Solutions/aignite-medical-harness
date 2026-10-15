"""SRRSH ICC 临床表与多模态文件夹 → AMA Dataset。

映射全部硬编码，不调用 LLM。icc-xlsx 每例固定输出 5 条临床表 Evidence；
icc-multimodal 复用这 5 条映射，并追加 MRI 报告与可用的 NIfTI/SVS 目录
Evidence。NIfTI/SVS 像素不进入 runner，结局和评分 target 不导入。

ama-dataset 契约对 Evidence/Turn 禁止额外字段，因此可审计信息不放数据集本体：
- 逐条 Evidence 的 event_time 三件套 / 缺失清单 → provenance.jsonl（按 episode 一行）
- 源 hash / 疑值 / 时间代理规则 / 缺失统计 → import_report.json
两者均不被 run loader 读取。
"""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..data import validate_dataset
from ..dataset_card import write_dataset_cards

SHEET_NAME = "ICC临床数据(合并版)"
HEADER_ROW_INDEX = 1
SOURCE_LOCATOR = "ICC临床数据.xlsx#主表!ID={id}"
TZ = timezone(timedelta(hours=8))

CLIN_FIELDS = [
    "Gender", "Age", "height", "weight", "BMI", "smoke", "alcohol",
    "HBV", "hypertension", "diabetes", "Cirrhosis", "child_pugh_grade",
]
LAB_FIELDS = ["AFP", "AST", "ALP", "GGT", "ALB", "BIL", "PT", "PLT", "INR"]
PATH_FIELDS = [
    "T", "N", "M", "Tumor_size", "tumor_large_vascular_invasion",
    "Tumor_number", "nerve_invasion", "MVI", "Grade",
]
REQUIRED_COLUMNS = set(
    CLIN_FIELDS + LAB_FIELDS + PATH_FIELDS
    + ["ID", "surg_time", "pathology", "TACE", "TACE_次数"]
)

ORIGIN_TABLE1 = "临床数据总表"
ORIGIN_TABLE2 = "表2"
TURN_ID = "t1"

CLINICAL_INSTRUCTIONS_TEXT = """本数据集用于无评分的证据检查。
仅依据当前轮已提供的 Evidence，在 answer 中返回 {"summary": "简明、基于证据的总结"}。
"""
MULTIMODAL_INSTRUCTIONS_TEXT = CLINICAL_INSTRUCTIONS_TEXT + (
    "NIfTI 和 SVS Evidence 仅登记原始文件的相对路径，模型未读取其像素内容；"
    "不得据此推断影像或病理所见。\n"
)

GENDER_TEXT = {"Male": "男性", "Female": "女性"}
YES_NO_TEXT = {
    "smoke": {"Yes": "是", "No": "否"},
    "alcohol": {"Yes": "是", "No": "否"},
    "HBV": {"Yes": "阳性", "No": "阴性"},
    "hypertension": {"Yes": "有", "No": "无"},
    "diabetes": {"Yes": "有", "No": "无"},
    "Cirrhosis": {"Yes": "有", "No": "无"},
    "tumor_large_vascular_invasion": {"Yes": "有", "No": "无"},
    "nerve_invasion": {"Yes": "有", "No": "无"},
    "MVI": {"Yes": "有", "No": "无"},
}
GRADE_TEXT = {
    "Low": "低级别",
    "Low_medium": "低-中级别",
    "Medium": "中级别",
    "Medium_high": "中-高级别",
    "High": "高级别",
}
CHILD_PUGH_TEXT = {
    "A": "Child-Pugh 分级 A 级（肝功能储备评级）",
    "B": "Child-Pugh 分级 B 级（肝功能储备评级）",
}

MISSING_TEXT = "未提供"
NOT_EVALUATED_MARK = "（未评估，非阴性）"
BASIS_SOURCE_RECORDED = "source_recorded"
BASIS_DERIVED_SPECIMEN = "derived_specimen_from_surg_time"
BASIS_PROXY_PREOP = "proxy_preop_anchor_surg_day"
NOTE_CLIN = "状态型资料（年龄/习惯/合并症），截至手术日有效；具体采集日期源数据未提供"
NOTE_LAB = "真实采样日期源数据未提供"
NOTE_TACE_NEG = "阴性治疗史：'截至手术日未行 TACE' 有效，非缺失"
NOTE_PATH = "取材=手术日（由 surg_time 推导，真实）"
UNITS_NOTE_CLIN = "单位确定：身高 cm、体重 kg。"
DIAGNOSIS_CN = "肝内胆管细胞癌（ICC）"

DATASET_LICENSE = "内部研究使用；含患者相关临床数据，不可再分发；对外发布前必须重编码"

TIME_PROXY_NOTE = (
    "唯一真实时间 = surg_time（treatment 依据 source_recorded）；基线/检验/TACE 的 "
    "event_time 为手术日代理（proxy_preop_anchor_surg_day）；病理 event_time 为取材推导"
    "（derived_specimen_from_surg_time）。逐条三件套见 provenance.jsonl"
)

# MRI 报告表（ICC核磁影像报告.xlsx）join 用常量：单 sheet、表头在第 0 行、无检查日期列。
MRI_SHEET_NAME = "sheet1"
MRI_FINDINGS_COLUMN = "检查所见"
MRI_DIAGNOSIS_COLUMN = "诊断"
MRI_SOURCE_LOCATOR = "影像数据/ICC核磁影像报告.xlsx#sheet1!住院号={id}"
MRI_EVIDENCE_ID = "mri-report-1"
MRI_EVIDENCE_TYPE = "radiology_report"
MRI_REPORT_PREFIX = "腹部 MRI 检查报告全文（检查日期源数据未提供，按术前检查解释）。"
MRI_TIME_NOTE = "MRI 检查日期源数据未提供；仅标记为术前相对时间，不生成精确 event_time"
# 术后/随访语境关键词：命中只在 provenance 打标，不改文本。
POSTOP_SUSPICION_KEYWORDS = ["术后", "复发", "随访", "复查", "TACE后", "介入后"]


def _fmt(value: Any) -> str:
    """将 Excel 数字格式化为稳定文本，整数值不保留 .0。"""
    if isinstance(value, bool):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value)


def _value(value: Any, mapping: dict[str, str] | None = None) -> str:
    if value is None:
        return MISSING_TEXT
    if mapping is not None:
        return mapping.get(str(value), _fmt(value))
    return _fmt(value)


def _table_origin(row: dict[str, Any]) -> str:
    """T 非空的记录来自临床数据总表；否则来自合并表中的表2。"""
    return ORIGIN_TABLE1 if row.get("T") is not None else ORIGIN_TABLE2


def _parse_surg_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()


def _event_time(surg_date: date) -> str:
    return datetime.combine(surg_date, time.min, tzinfo=TZ).isoformat()


def _missing(row: dict[str, Any], fields: list[str]) -> list[str]:
    return [field for field in fields if row.get(field) is None]


def build_clinical_baseline(row: dict[str, Any], ctx: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """返回 (evidence, provenance 片段)。"""
    age = row.get("Age")
    gender = _value(row.get("Gender"), GENDER_TEXT)
    if age is None:
        head = f"年龄{MISSING_TEXT}，{gender}"
    else:
        head = f"{_fmt(age)} 岁{gender}"

    text = (
        f"{head}。身高 {_value(row.get('height'))} cm，体重 {_value(row.get('weight'))} kg，"
        f"BMI {_value(row.get('BMI'))}。"
        f"吸烟：{_value(row.get('smoke'), YES_NO_TEXT['smoke'])}；"
        f"饮酒：{_value(row.get('alcohol'), YES_NO_TEXT['alcohol'])}。"
        f"HBV（乙肝病毒感染）：{_value(row.get('HBV'), YES_NO_TEXT['HBV'])}。"
        f"高血压：{_value(row.get('hypertension'), YES_NO_TEXT['hypertension'])}；"
        f"糖尿病：{_value(row.get('diabetes'), YES_NO_TEXT['diabetes'])}；"
        f"肝硬化：{_value(row.get('Cirrhosis'), YES_NO_TEXT['Cirrhosis'])}。"
    )
    child_pugh = row.get("child_pugh_grade")
    if child_pugh is None:
        text += f"Child-Pugh 分级：{MISSING_TEXT}。"
    else:
        text += CHILD_PUGH_TEXT.get(
            str(child_pugh), f"Child-Pugh 分级 {_fmt(child_pugh)}"
        ) + "。"

    provenance = {
        "event_time": ctx["surg_iso"],
        "event_time_basis": BASIS_PROXY_PREOP,
        "event_time_note": NOTE_CLIN,
        "table_origin": _table_origin(row),
        "units_note": UNITS_NOTE_CLIN,
    }
    missing = _missing(row, CLIN_FIELDS)
    if missing:
        provenance["missing_fields"] = missing
    return {
        "id": "clin-1",
        "type": "clinical_baseline",
        "text": text,
    }, provenance


def build_lab_panel(row: dict[str, Any], ctx: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    values = "；".join(f"{field} {_value(row.get(field))}" for field in LAB_FIELDS)
    provenance = {
        "event_time": ctx["surg_iso"],
        "event_time_basis": BASIS_PROXY_PREOP,
        "event_time_note": NOTE_LAB,
    }
    missing = _missing(row, LAB_FIELDS)
    if missing:
        provenance["missing_fields"] = missing
    return {
        "id": "lab-1",
        "type": "lab_panel",
        "text": (
            "围术期血清学检验 panel（采样时点源数据未提供，按术前基线解释；"
            f"单位未随源数据提供）：{values}。"
        ),
    }, provenance


def build_tace(row: dict[str, Any], ctx: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    tace = row.get("TACE")
    count = row.get("TACE_次数")
    if str(tace) == "No":
        text = "术前未行 TACE（肝动脉化疗栓塞，术前辅助介入治疗）。"
        note = NOTE_TACE_NEG
    elif str(tace) == "Yes":
        count_text = "次数未提供" if count is None else f"次数 {_fmt(count)} 次"
        text = f"术前曾行 TACE（肝动脉化疗栓塞，术前辅助介入治疗），{count_text}。"
        note = "阳性治疗史：截至手术日曾行 TACE"
    elif tace is None:
        text = f"术前 TACE 治疗史：{MISSING_TEXT}。"
        note = "TACE 治疗史源数据未提供"
    else:
        text = f"术前 TACE 治疗史：{_fmt(tace)}。"
        note = "TACE 治疗史按源值保留"

    provenance = {
        "event_time": ctx["surg_iso"],
        "event_time_basis": BASIS_PROXY_PREOP,
        "event_time_note": note,
    }
    if tace is not None:
        provenance["tace"] = str(tace)
    missing = _missing(row, ["TACE", "TACE_次数"])
    if missing:
        provenance["missing_fields"] = missing
    return {
        "id": "tx-tace-1",
        "type": "treatment_record",
        "text": text,
    }, provenance


def build_mri_report(
    pid: str,
    report: dict[str, str],
) -> tuple[dict[str, Any], dict[str, Any]]:
    """MRI 报告全文 → (evidence, provenance 片段)。检查所见/诊断保持源文本原样。"""
    findings = report.get(MRI_FINDINGS_COLUMN) or MISSING_TEXT
    diagnosis = report.get(MRI_DIAGNOSIS_COLUMN) or MISSING_TEXT
    evidence = {
        "id": MRI_EVIDENCE_ID,
        "type": MRI_EVIDENCE_TYPE,
        "text": f"{MRI_REPORT_PREFIX}\n检查所见：{findings}\n诊断：{diagnosis}",
    }
    provenance = {
        "relative_time": "preoperative",
        "event_time_missing": True,
        "event_time_note": MRI_TIME_NOTE,
        "source": MRI_SOURCE_LOCATOR.format(id=pid),
        "fields": [MRI_FINDINGS_COLUMN, MRI_DIAGNOSIS_COLUMN],
    }
    # 术后/随访语境仅打标（按常量列表顺序），供评测侧甄别，绝不改写或剔除文本。
    suspected = [
        keyword for keyword in POSTOP_SUSPICION_KEYWORDS
        if keyword in f"{findings}{diagnosis}"
    ]
    if suspected:
        provenance["postop_suspicion_keywords"] = suspected
    return evidence, provenance


def build_surgery(row: dict[str, Any], ctx: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    return {
        "id": "tx-surg-1",
        "type": "treatment_record",
        "text": f"患者于 {ctx['surg_date'].isoformat()} 接受手术治疗（术式源数据未收录）。",
    }, {
        "event_time": ctx["surg_iso"],
        "event_time_basis": BASIS_SOURCE_RECORDED,
        "missing_fields": ["术式"],
    }


def _stage(value: Any, label: str) -> str:
    if value is None:
        return f"{label}：{MISSING_TEXT}"
    value_text = str(value)
    if value_text.lower() == f"{label.lower()}x":
        return f"{label}：{value_text}{NOT_EVALUATED_MARK}"
    return f"{label}：{value_text}"


def build_pathology(row: dict[str, Any], ctx: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    pathology = row.get("pathology")
    diagnosis = DIAGNOSIS_CN if str(pathology) == "ICC" else _value(pathology)
    grade = row.get("Grade")
    if grade is None:
        grade_text = MISSING_TEXT
    elif str(grade) in GRADE_TEXT:
        grade_text = f"{GRADE_TEXT[str(grade)]}（Grade: {grade}）"
    else:
        grade_text = _fmt(grade)

    t_stage = row.get("T")
    n_stage = row.get("N")
    m_stage = row.get("M")
    text = (
        f"术后病理：{diagnosis}。"
        f"肿瘤最大径 {_value(row.get('Tumor_size'))} cm；"
        f"肿瘤数目 {_value(row.get('Tumor_number'))}；"
        "大血管侵犯："
        f"{_value(row.get('tumor_large_vascular_invasion'), YES_NO_TEXT['tumor_large_vascular_invasion'])}；"
        f"神经侵犯：{_value(row.get('nerve_invasion'), YES_NO_TEXT['nerve_invasion'])}；"
        f"MVI（微血管侵犯）：{_value(row.get('MVI'), YES_NO_TEXT['MVI'])}；"
        f"组织学分级：{grade_text}。"
        f"TNM 分期：{_value(t_stage)}；{_stage(n_stage, 'N')}；{_stage(m_stage, 'M')}。"
    )
    provenance = {
        "event_time": ctx["surg_iso"],
        "event_time_basis": BASIS_DERIVED_SPECIMEN,
        "event_time_note": NOTE_PATH,
    }
    missing = _missing(row, PATH_FIELDS)
    if missing:
        provenance["missing_fields"] = missing
    not_evaluated = [
        label for label, value in (("N", n_stage), ("M", m_stage))
        if value is not None and str(value).lower() == f"{label.lower()}x"
    ]
    if not_evaluated:
        provenance["not_evaluated"] = not_evaluated
    return {
        "id": "path-1",
        "type": "pathology_registry",
        "text": text,
    }, provenance


def build_episode(row: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    """一行患者 → 5 条临床 Evidence 的 (episode, provenance 行)。"""
    pid = _fmt(row["ID"])
    surg_date = _parse_surg_date(row["surg_time"])
    ctx = {
        "surg_date": surg_date,
        "surg_iso": _event_time(surg_date),
        "locator": SOURCE_LOCATOR.format(id=pid),
    }
    built = [
        build_clinical_baseline(row, ctx),
        build_lab_panel(row, ctx),
        build_tace(row, ctx),
        build_surgery(row, ctx),
        build_pathology(row, ctx),
    ]
    evidence = [item[0] for item in built]
    episode = {
        "id": f"icc-{pid}",
        "turns": [{"id": TURN_ID, "evidence": evidence}],
    }
    provenance_row = {
        "id": episode["id"],
        "source": ctx["locator"],
        "surg_time": surg_date.isoformat(),
        "table_origin": _table_origin(row),
        "evidence": {item[0]["id"]: item[1] for item in built},
    }
    return episode, provenance_row


def _read_main_table(source: Path) -> list[dict[str, Any]]:
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise FileNotFoundError(f"sheet '{SHEET_NAME}' not in {source} (have {wb.sheetnames})")
        rows = list(wb[SHEET_NAME].iter_rows(values_only=True))
    finally:
        wb.close()

    header = [str(value).strip() if value is not None else "" for value in rows[HEADER_ROW_INDEX]]
    missing_columns = sorted(REQUIRED_COLUMNS - set(header))
    if missing_columns:
        raise ValueError(f"主表缺少必需列: {missing_columns}")
    data = [
        dict(zip(header, values))
        for values in rows[HEADER_ROW_INDEX + 1:]
        if any(value is not None for value in values)
    ]
    return data


def import_srrsh_icc(
    source: Path,
    out: Path,
    ids: list[str] | None = None,
) -> dict[str, Any]:
    """把 ICC 临床主表导入为固定 5 条 Evidence 的无 target 数据集。"""
    source = Path(source)
    out = Path(out)
    rows = _read_main_table(source)
    out.mkdir(parents=True, exist_ok=True)

    wanted = {str(value).strip() for value in ids} if ids else None
    episodes: list[dict[str, Any]] = []
    provenance: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for row in sorted(rows, key=lambda item: _fmt(item["ID"])):
        pid = _fmt(row["ID"])
        if wanted is not None and pid not in wanted:
            continue
        if pid in seen_ids:
            skipped.append({"id": pid, "reason": "duplicate ID"})
            continue
        try:
            episode, provenance_row = build_episode(row)
        except (KeyError, TypeError, ValueError) as exc:
            skipped.append({"id": pid, "reason": f"{type(exc).__name__}: {exc}"})
            continue
        seen_ids.add(pid)
        episodes.append(episode)
        provenance.append(provenance_row)
    episode_ids = [episode["id"] for episode in episodes]
    (out / "dataset.json").write_text(
        json.dumps({"schema": "ama-dataset", "name": out.name,
                    "splits": {"all": episode_ids}}, indent=2),
        encoding="utf-8",
    )
    (out / "episodes.jsonl").write_text(
        "".join(json.dumps(episode, ensure_ascii=False) + "\n" for episode in episodes),
        encoding="utf-8",
    )
    (out / "provenance.jsonl").write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in provenance),
        encoding="utf-8",
    )

    tracked_columns = CLIN_FIELDS + LAB_FIELDS + PATH_FIELDS + ["TACE", "TACE_次数"]
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "source_file": str(source),
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "sheet": SHEET_NAME,
        "n_rows": len(rows),
        "n_imported": len(episodes),
        "skipped": skipped,
        "ids_filter": sorted(wanted) if wanted is not None else None,
        "table_origin": {
            "rule": "T 列非空 → 临床数据总表（表1，188 例）；否则表2（68 例）",
            "counts": {
                ORIGIN_TABLE1: sum(1 for p in provenance if p["table_origin"] == ORIGIN_TABLE1),
                ORIGIN_TABLE2: sum(1 for p in provenance if p["table_origin"] == ORIGIN_TABLE2),
            },
        },
        "time_proxy": TIME_PROXY_NOTE,
        "outcome_columns": {
            "policy": "RFS/OS/随访列不读取、不进 episodes/provenance（评测侧负责 targets）",
        },
        "missing_stats": {
            "n_rows": len(rows),
            "nonnull_by_column": {c: sum(1 for r in rows if r.get(c) is not None)
                                  for c in tracked_columns},
        },
        "counts": {
            "episodes": len(episodes),
            "turns": len(episodes),
            "evidence": 5 * len(episodes),
        },
    }
    (out / "import_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (out / "instructions.txt").write_text(CLINICAL_INSTRUCTIONS_TEXT, encoding="utf-8")
    write_dataset_cards(
        out,
        name=out.name,
        source=f"{source}#{SHEET_NAME}",
        license_name=DATASET_LICENSE,
        purpose_en=("SRRSH ICC clinical-table import: one single-turn episode per patient with "
                    "5 table-derived evidence items."),
        purpose_zh="邵逸夫 ICC 临床表导入：每位患者 1 个单轮 Episode、5 条临床表 Evidence。",
        construction_en=("Hard-coded column mapping, no LLM. Outcome columns are never read. "
                         "Per-evidence timing and missingness are stored in provenance.jsonl, "
                         "which the run loader never reads."),
        construction_zh=("列到 Evidence 的映射硬编码，无 LLM；结局列不读取。逐条时间与缺失信息"
                         "保存在 provenance.jsonl，run loader 不读取。"),
        episodes=len(episodes), turns=len(episodes),
        evidence=5 * len(episodes),
        scorer="unscored", targets=0,
        limitations_en=("Time anchors are declared proxies (only surg_time is real). Lab values "
                        "carry no units. Output contains patient identifiers: keep it out of the "
                        "repository and do not redistribute before re-encoding."),
        limitations_zh=("时间锚为声明的代理（唯一真实时间为 surg_time）；检验值不带单位。产物含"
                        "患者标识：不得进入仓库，再分发前必须重编码。"),
    )
    return report


# ---------------------------------------------------------------- multimodal folder importer

CLINICAL_XLSX = Path("ICC临床数据.xlsx")
REPORT_XLSX = Path("影像数据/ICC核磁影像报告.xlsx")
MRI_DIR = Path("影像数据/初期182例ICC_MRI")
WSI_DIR = Path("病理数据")
REPORT_SHEET = MRI_SHEET_NAME
REPORT_COLUMNS = ("住院号", "检查所见", "诊断")

MRI_SEQUENCES = ("DWI", "T1_ART", "T1_DL", "T1_NC", "T1_PV", "T2")
MRI_VARIANTS = ("raw", "n4", "roi")
MRI_VARIANT_LABELS = {"raw": "原始", "n4": "N4 偏置场校正", "roi": "ROI 分割掩膜"}
EXPECTED_MRI_KEYS = tuple(
    (sequence, variant)
    for sequence in MRI_SEQUENCES
    for variant in MRI_VARIANTS
)


@dataclass
class _MultimodalInventory:
    root: Path
    clinical_rows: dict[str, dict[str, Any]]
    selected_ids: list[str]
    reports: dict[str, dict[str, str]]
    mri: dict[str, dict[tuple[str, str], Path]]
    wsi: dict[str, list[Path]]
    ignored_files: list[str]
    mri_directory_exists: bool
    wsi_directory_exists: bool


def _patient_id(value: Any) -> str:
    if value is None:
        raise ValueError("患者 ID 不能为空")
    patient_id = _fmt(value).strip()
    if not patient_id:
        raise ValueError("患者 ID 不能为空")
    return patient_id


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _is_ignored(path: Path) -> bool:
    return (
        path.name == ".DS_Store"
        or path.name.startswith("._")
        or path.suffix.lower() in {".zip", ".dcm"}
    )


def _relative_path(root: Path, path: Path) -> str:
    root_resolved = root.resolve()
    path_resolved = path.resolve()
    if path_resolved != root_resolved and root_resolved not in path_resolved.parents:
        raise ValueError(f"artifact 路径逃逸数据集根目录: {path}")
    try:
        return path.relative_to(root).as_posix()
    except ValueError as exc:
        raise ValueError(f"artifact 不在数据集根目录内: {path}") from exc


def _read_clinical_rows(path: Path) -> dict[str, dict[str, Any]]:
    indexed: dict[str, dict[str, Any]] = {}
    for row in _read_main_table(path):
        patient_id = _patient_id(row.get("ID"))
        if patient_id in indexed:
            raise ValueError(f"临床主表存在重复 ID: {patient_id}")
        indexed[patient_id] = row
    return indexed


def _read_reports(path: Path) -> dict[str, dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if REPORT_SHEET not in workbook.sheetnames:
            raise FileNotFoundError(
                f"sheet '{REPORT_SHEET}' not in {path} (have {workbook.sheetnames})"
            )
        rows = workbook[REPORT_SHEET].iter_rows(values_only=True)
        try:
            header_values = next(rows)
        except StopIteration as exc:
            raise ValueError(f"影像报告表为空: {path}") from exc
        header = [str(value).strip() if value is not None else "" for value in header_values]
        missing = sorted(set(REPORT_COLUMNS) - set(header))
        if missing:
            raise ValueError(f"影像报告表缺少必需列: {missing}")

        indexed: dict[str, dict[str, str]] = {}
        for values in rows:
            if not any(value is not None for value in values):
                continue
            row = dict(zip(header, values))
            patient_id = _patient_id(row.get("住院号"))
            if patient_id in indexed:
                raise ValueError(f"影像报告表存在重复住院号: {patient_id}")
            indexed[patient_id] = {
                "findings": str(row.get("检查所见") or "").strip(),
                "diagnosis": str(row.get("诊断") or "").strip(),
            }
        return indexed
    finally:
        workbook.close()


def _mri_key(path: Path) -> tuple[str, str] | None:
    lowered = path.name.lower()
    if lowered.endswith(".nii.gz"):
        stem = path.name[:-7]
    elif lowered.endswith(".nii"):
        stem = path.name[:-4]
    else:
        return None
    expected = {sequence.upper(): (sequence, "raw") for sequence in MRI_SEQUENCES}
    expected.update({
        f"{sequence}_N4".upper(): (sequence, "n4")
        for sequence in MRI_SEQUENCES
    })
    expected.update({
        f"{sequence}_ROI".upper(): (sequence, "roi")
        for sequence in MRI_SEQUENCES
    })
    return expected.get(stem.upper())


def _directory_patient_id(name: str) -> str | None:
    match = re.fullmatch(r"([0-9]+)(?:_.+)?", name)
    return match.group(1) if match else None


def _scan_mri(
    root: Path,
    directory: Path,
) -> tuple[dict[str, dict[tuple[str, str], Path]], list[str]]:
    indexed: dict[str, dict[tuple[str, str], Path]] = {}
    ignored: list[str] = []
    for entry in sorted(directory.iterdir(), key=lambda item: item.name):
        if _is_ignored(entry) or not entry.is_dir():
            ignored.append(_relative_path(root, entry))
            continue
        patient_id = _directory_patient_id(entry.name)
        if patient_id is None:
            ignored.append(_relative_path(root, entry))
            continue
        if patient_id in indexed:
            raise ValueError(f"MRI 目录存在重复患者 ID: {patient_id}")
        artifacts: dict[tuple[str, str], Path] = {}
        for path in sorted(entry.iterdir(), key=lambda item: item.name):
            if _is_ignored(path) or not path.is_file():
                ignored.append(_relative_path(root, path))
                continue
            key = _mri_key(path)
            if key is None:
                ignored.append(_relative_path(root, path))
                continue
            if key in artifacts:
                raise ValueError(
                    f"MRI 患者 {patient_id} 存在重复序列文件: {key[0]} {key[1]}"
                )
            _relative_path(root, path)
            artifacts[key] = path
        indexed[patient_id] = artifacts
    return indexed, ignored


def _wsi_patient_id(path: Path) -> str | None:
    match = re.match(r"^([0-9]+)(?:-|$)", path.stem)
    return match.group(1) if match else None


def _scan_wsi(root: Path, directory: Path) -> tuple[dict[str, list[Path]], list[str]]:
    indexed: dict[str, list[Path]] = {}
    ignored: list[str] = []
    for path in sorted(directory.iterdir(), key=lambda item: item.name):
        if _is_ignored(path) or not path.is_file() or path.suffix.lower() != ".svs":
            ignored.append(_relative_path(root, path))
            continue
        patient_id = _wsi_patient_id(path)
        if patient_id is None:
            ignored.append(_relative_path(root, path))
            continue
        _relative_path(root, path)
        indexed.setdefault(patient_id, []).append(path)
    return indexed, ignored


def _collect_multimodal_inventory(
    source: Path,
    ids: list[str] | None,
) -> _MultimodalInventory:
    root = Path(source).expanduser().resolve()
    if not root.is_dir():
        raise FileNotFoundError(f"数据根目录不存在: {source}")
    required = [root / CLINICAL_XLSX, root / REPORT_XLSX]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"缺少必需输入: {missing}")

    mri_directory = root / MRI_DIR
    wsi_directory = root / WSI_DIR
    for path in (mri_directory, wsi_directory):
        if path.exists() and not path.is_dir():
            raise NotADirectoryError(f"可选模态路径存在但不是目录: {path}")
    mri_directory_exists = mri_directory.is_dir()
    wsi_directory_exists = wsi_directory.is_dir()

    clinical_rows = _read_clinical_rows(root / CLINICAL_XLSX)
    reports = _read_reports(root / REPORT_XLSX)
    mri, ignored_mri = _scan_mri(root, mri_directory) if mri_directory_exists else ({}, [])
    wsi, ignored_wsi = _scan_wsi(root, wsi_directory) if wsi_directory_exists else ({}, [])

    wanted = {_patient_id(value) for value in ids} if ids else None
    unknown = sorted(wanted - set(clinical_rows)) if wanted is not None else []
    if unknown:
        raise ValueError(f"指定 ID 不在临床主表中: {unknown}")
    selected_ids = sorted(wanted if wanted is not None else clinical_rows)
    return _MultimodalInventory(
        root=root,
        clinical_rows=clinical_rows,
        selected_ids=selected_ids,
        reports=reports,
        mri=mri,
        wsi=wsi,
        ignored_files=sorted(ignored_mri + ignored_wsi),
        mri_directory_exists=mri_directory_exists,
        wsi_directory_exists=wsi_directory_exists,
    )


def _artifact_record(root: Path, path: Path) -> dict[str, Any]:
    stat = path.stat()
    return {
        "path": _relative_path(root, path),
        "size_bytes": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
    }


def _multimodal_audit_payload(
    inventory: _MultimodalInventory,
    ids: list[str] | None,
) -> dict[str, Any]:
    clinical_ids = set(inventory.clinical_rows)
    report_ids = set(inventory.reports)
    mri_ids = set(inventory.mri)
    wsi_ids = set(inventory.wsi)
    selected = set(inventory.selected_ids)
    empty_reports = sorted(
        patient_id
        for patient_id, report in inventory.reports.items()
        if not report["findings"] and not report["diagnosis"]
    )
    mri_missing_files = {
        patient_id: [
            f"{sequence}:{variant}"
            for sequence, variant in EXPECTED_MRI_KEYS
            if (sequence, variant) not in inventory.mri[patient_id]
        ]
        for patient_id in sorted(selected & mri_ids)
    }
    mri_missing_files = {
        patient_id: missing
        for patient_id, missing in mri_missing_files.items()
        if missing
    }
    artifact_paths = [
        path
        for patient_id in inventory.selected_ids
        for path in (
            list(inventory.mri.get(patient_id, {}).values())
            + inventory.wsi.get(patient_id, [])
        )
    ]
    return {
        "source_root": str(inventory.root),
        "inputs": {
            "clinical_xlsx": {
                "path": CLINICAL_XLSX.as_posix(),
                "sha256": _sha256(inventory.root / CLINICAL_XLSX),
            },
            "radiology_report_xlsx": {
                "path": REPORT_XLSX.as_posix(),
                "sha256": _sha256(inventory.root / REPORT_XLSX),
            },
            "mri_directory": {
                "path": MRI_DIR.as_posix(),
                "exists": inventory.mri_directory_exists,
            },
            "wsi_directory": {
                "path": WSI_DIR.as_posix(),
                "exists": inventory.wsi_directory_exists,
            },
        },
        "ids_filter": sorted({_patient_id(value) for value in ids}) if ids else None,
        "cohort": {
            "clinical_patients": len(clinical_ids),
            "report_patients": len(report_ids),
            "mri_patients": len(mri_ids),
            "wsi_patients": len(wsi_ids),
            "selected_patients": len(selected),
        },
        "coverage": {
            "selected_with_report": len(selected & report_ids),
            "selected_with_mri": len(selected & mri_ids),
            "selected_with_wsi": len(selected & wsi_ids),
            "clinical_with_all_modalities": len(clinical_ids & report_ids & mri_ids & wsi_ids),
        },
        "missing": {
            "selected_without_report": sorted(selected - report_ids),
            "selected_without_mri": sorted(selected - mri_ids),
            "selected_without_wsi": sorted(selected - wsi_ids),
            "empty_report_ids": empty_reports,
            "mri_expected_files": mri_missing_files,
            "missing_source_directories": [
                path.as_posix()
                for path, exists in (
                    (MRI_DIR, inventory.mri_directory_exists),
                    (WSI_DIR, inventory.wsi_directory_exists),
                )
                if not exists
            ],
        },
        "orphans": {
            "report_without_clinical": sorted(report_ids - clinical_ids),
            "mri_without_clinical": sorted(mri_ids - clinical_ids),
            "wsi_without_clinical": sorted(wsi_ids - clinical_ids),
        },
        "ignored_files": inventory.ignored_files,
        "evidence_policy": {
            "radiology_report_id": MRI_EVIDENCE_ID,
            "radiology_report_time": "preoperative; exact event_time unavailable",
            "binary_artifacts": (
                "NIfTI/SVS are catalogued by safe relative path; pixels are not loaded by ama run"
            ),
        },
        "artifacts": [
            _artifact_record(inventory.root, path)
            for path in sorted(artifact_paths, key=lambda item: item.as_posix())
        ],
    }


def audit_srrsh_icc_multimodal(
    source: Path,
    ids: list[str] | None = None,
) -> dict[str, Any]:
    """只读扫描 ICC 多模态数据根目录并返回覆盖与缺失报告。"""
    inventory = _collect_multimodal_inventory(source, ids)
    return _multimodal_audit_payload(inventory, ids)


def _report_evidence(
    patient_id: str,
    report: dict[str, str],
) -> tuple[dict[str, Any], dict[str, Any]] | None:
    if not report["findings"] and not report["diagnosis"]:
        return None
    return build_mri_report(patient_id, {
        MRI_FINDINGS_COLUMN: report["findings"],
        MRI_DIAGNOSIS_COLUMN: report["diagnosis"],
    })


def _mri_evidence(
    root: Path,
    artifacts: dict[tuple[str, str], Path],
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    evidence: list[dict[str, Any]] = []
    provenance: dict[str, dict[str, Any]] = {}
    for sequence, variant in EXPECTED_MRI_KEYS:
        path = artifacts.get((sequence, variant))
        if path is None:
            continue
        relative_path = _relative_path(root, path)
        sequence_id = sequence.lower().replace("_", "-")
        is_mask = variant == "roi"
        if is_mask:
            text = (
                f"MRI NIfTI 分割掩膜：{sequence} 序列 ROI；"
                f"原始文件：{relative_path}；像素内容未转写为文本。"
            )
        else:
            text = (
                f"MRI NIfTI 文件：{sequence} 序列，{MRI_VARIANT_LABELS[variant]}版本；"
                f"原始文件：{relative_path}；像素内容未转写为文本。"
            )
        evidence_id = f"mri-{sequence_id}-{variant}"
        evidence.append({
            "id": evidence_id,
            "type": "image_mask" if is_mask else "medical_image",
            "text": text,
        })
        provenance[evidence_id] = {
            "source": relative_path,
            "modality": "MRI",
            "sequence": sequence,
            "variant": variant,
            "format": "NIfTI",
            "relative_time": "preoperative",
            "event_time_missing": True,
        }
    return evidence, provenance


def _slug(value: str) -> str:
    slug = re.sub(r"[^0-9A-Za-z]+", "-", value).strip("-").lower()
    return slug or "slide"


def _wsi_evidence(
    root: Path,
    paths: list[Path],
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    evidence: list[dict[str, Any]] = []
    provenance: dict[str, dict[str, Any]] = {}
    used_ids: set[str] = set()
    for path in sorted(paths, key=lambda item: item.name):
        evidence_id = f"wsi-{_slug(path.stem)}"
        if evidence_id in used_ids:
            raise ValueError(f"WSI 文件名规范化后 Evidence ID 重复: {evidence_id}")
        used_ids.add(evidence_id)
        relative_path = _relative_path(root, path)
        evidence.append({
            "id": evidence_id,
            "type": "whole_slide_image",
            "text": (
                "病理全切片 SVS 文件；"
                f"原始文件：{relative_path}；切片像素内容未转写为文本。"
            ),
        })
        provenance[evidence_id] = {
            "source": relative_path,
            "modality": "pathology",
            "format": "SVS",
            "relative_time": "postoperative_pathology",
            "event_time_missing": True,
        }
    return evidence, provenance


def build_multimodal_episode(source: Path, patient_id: str) -> dict[str, Any]:
    """从多模态数据根目录只读构建一个指定患者 Episode。"""
    inventory = _collect_multimodal_inventory(source, [patient_id])
    episode, _ = _build_multimodal_episode(inventory, _patient_id(patient_id))
    return episode


def _build_multimodal_episode(
    inventory: _MultimodalInventory,
    patient_id: str,
) -> tuple[dict[str, Any], dict[str, Any]]:
    episode, provenance = build_episode(inventory.clinical_rows[patient_id])
    turn = episode["turns"][0]
    report = inventory.reports.get(patient_id)
    if report is not None:
        report_item = _report_evidence(patient_id, report)
        if report_item is not None:
            report_evidence, report_provenance = report_item
            turn["evidence"].insert(3, report_evidence)
            provenance["evidence"][report_evidence["id"]] = report_provenance
    mri_evidence, mri_provenance = _mri_evidence(
        inventory.root, inventory.mri.get(patient_id, {})
    )
    wsi_evidence, wsi_provenance = _wsi_evidence(
        inventory.root, inventory.wsi.get(patient_id, [])
    )
    turn["evidence"].extend(mri_evidence)
    turn["evidence"].extend(wsi_evidence)
    provenance["evidence"].update(mri_provenance)
    provenance["evidence"].update(wsi_provenance)
    return episode, provenance


def import_srrsh_icc_multimodal(
    source: Path,
    out: Path,
    ids: list[str] | None = None,
    dataset_name: str = "srrsh-icc-multimodal",
) -> dict[str, Any]:
    """在 ICC NAS 根目录原位生成无 target 的 AMA Dataset。"""
    source_root = Path(source).expanduser().resolve()
    out_root = Path(out).expanduser().resolve()
    if source_root != out_root:
        raise ValueError("icc-multimodal 要求 --source 与 --out 指向同一数据根目录")
    if (source_root / "targets.jsonl").exists():
        raise FileExistsError("数据根目录已存在 targets.jsonl，为避免混入评分标签已停止导入")

    inventory = _collect_multimodal_inventory(source_root, ids)
    built = [
        _build_multimodal_episode(inventory, patient_id)
        for patient_id in inventory.selected_ids
    ]
    episodes = [item[0] for item in built]
    provenance = [item[1] for item in built]
    episode_ids = [episode["id"] for episode in episodes]
    dataset = {
        "schema": "ama-dataset",
        "name": dataset_name,
        "splits": {"all": episode_ids},
    }
    (source_root / "dataset.json").write_text(
        json.dumps(dataset, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (source_root / "episodes.jsonl").write_text(
        "".join(json.dumps(episode, ensure_ascii=False) + "\n" for episode in episodes),
        encoding="utf-8",
    )
    (source_root / "provenance.jsonl").write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in provenance),
        encoding="utf-8",
    )

    counts_by_type: dict[str, int] = {}
    for episode in episodes:
        for evidence in episode["turns"][0]["evidence"]:
            evidence_type = evidence["type"]
            counts_by_type[evidence_type] = counts_by_type.get(evidence_type, 0) + 1
    report = _multimodal_audit_payload(inventory, ids)
    report.update({
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "dataset_name": dataset_name,
        "counts": {
            "episodes": len(episodes),
            "turns": len(episodes),
            "evidence": sum(counts_by_type.values()),
            "evidence_by_type": counts_by_type,
        },
    })
    (source_root / "import_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (source_root / "instructions.txt").write_text(
        MULTIMODAL_INSTRUCTIONS_TEXT, encoding="utf-8"
    )
    write_dataset_cards(
        source_root,
        name=dataset_name,
        source=(f"{CLINICAL_XLSX.as_posix()}; {REPORT_XLSX.as_posix()}; optional "
                f"{MRI_DIR.as_posix()} and {WSI_DIR.as_posix()}"),
        license_name=DATASET_LICENSE,
        purpose_en=("SRRSH ICC import with clinical evidence and radiology reports, plus available "
                    "MRI NIfTI and pathology SVS catalog entries."),
        purpose_zh=("邵逸夫 ICC 临床与 MRI 报告导入，并按实际可用情况登记 MRI NIfTI 和病理 "
                    "SVS 文件目录。"),
        construction_en=("Hard-coded mapping, no LLM. The clinical workbook defines the cohort. "
                         "MRI and WSI directories are optional; missing modalities do not remove "
                         "patients. Evidence provenance is stored in provenance.jsonl and is not "
                         "read by the runner."),
        construction_zh=("硬编码映射，无 LLM；临床表定义患者队列。MRI 与 WSI 目录可选，缺失模态"
                         "不剔除患者。逐条 Evidence 的溯源信息保存在 provenance.jsonl，runner "
                         "不读取该文件。"),
        episodes=len(episodes),
        turns=len(episodes),
        evidence=sum(counts_by_type.values()),
        scorer="unscored",
        targets=0,
        artifacts_en=("NIfTI and SVS files remain in place below the dataset root and are not "
                      "copied. Evidence text records safe relative paths; ama run does not load "
                      "their pixels. Paths, sizes and modification times are audited in "
                      "import_report.json."),
        artifacts_zh=("NIfTI 与 SVS 保留在数据集根目录原位，不复制。Evidence 文本仅登记安全相对"
                      "路径，ama run 不读取其像素；路径、大小和修改时间记录在 import_report.json。"),
        limitations_en=("Output contains patient identifiers and must never enter the repository "
                        "before de-identification. Radiology reports have no exact examination "
                        "date. Catalog entries do not expose image or slide pixels to the model."),
        limitations_zh=("产物含患者标识，脱敏前不得进入仓库。MRI 报告无准确检查日期；NIfTI/SVS "
                        "目录 Evidence 不向模型提供影像或切片像素。"),
    )

    errors = validate_dataset(source_root)
    if errors:
        raise ValueError(f"AMA Dataset 校验失败: {errors}")
    return report


def configure_clinical_cli(parser: Any) -> None:
    parser.add_argument("--id", dest="ids", action="append")


def import_clinical_from_cli(args: Any) -> dict[str, Any]:
    return import_srrsh_icc(args.source, args.out, ids=args.ids)


def configure_multimodal_cli(parser: Any) -> None:
    parser.add_argument("--id", dest="ids", action="append")
    parser.add_argument("--name", default="srrsh-icc-multimodal")


def import_multimodal_from_cli(args: Any) -> dict[str, Any]:
    return import_srrsh_icc_multimodal(
        args.source,
        args.out,
        ids=args.ids,
        dataset_name=args.name,
    )
